from __future__ import annotations

import csv
import io
import json
import math
import re
from collections import defaultdict
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle

NAVY = HexColor("#0e2138")
INK = HexColor("#17243a")
MUTED = HexColor("#68778a")
GRID = HexColor("#d8e0ea")
LIGHT = HexColor("#eef3f8")
RED_BG = HexColor("#fff1ef")
RED = HexColor("#b52b22")
BLUE = HexColor("#2f86c1")
PAGE_W, PAGE_H = A4

SERIES_META = {
    "CUSA": ("Coimbra", "Usado"),
    "CNOV": ("Coimbra", "Novo"),
    "PUSA": ("Picoto", "Usado"),
    "PNOV": ("Picoto", "Novo"),
    "POFI": ("Picoto", "Oficina"),
}


def money(value: float) -> str:
    value = float(value or 0)
    s = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} EUR"


def pct(value: float) -> str:
    return f"{value:.1f}%".replace(".", ",")


def as_float(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("€", "").replace("EUR", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def normalize_day(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value[:10], fmt).date().isoformat()
            except ValueError:
                pass
    return None


def parse_multipart(headers, body: bytes):
    content_type = headers.get("Content-Type", "")
    if "multipart/form-data" not in content_type:
        raise ValueError("O relatório requer data e Excel dos vendedores.")
    raw = (
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8")
        + body
    )
    message = BytesParser(policy=policy.default).parsebytes(raw)
    fields: dict[str, str] = {}
    files: dict[str, tuple[str, bytes]] = {}
    for part in message.iter_parts():
        params = dict(part.get_params(header="content-disposition") or [])
        name = params.get("name")
        if not name:
            continue
        payload = part.get_payload(decode=True) or b""
        filename = params.get("filename")
        if filename:
            files[name] = (filename, payload)
        else:
            fields[name] = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
    return fields, files


def parse_seller_workbook(filename: str, content: bytes, report_date: str):
    result = {
        "daily": defaultdict(float),
        "monthly": defaultdict(float),
        "daily_by_location": defaultdict(float),
        "monthly_by_location": defaultdict(float),
        "daily_series": defaultdict(float),
        "scrap_day": 0.0,
        "scrap_month": 0.0,
        "source": filename,
    }

    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            seller = (row.get("Vendedor") or row.get("vendedor") or row.get("Nome") or "").strip()
            day = normalize_day(row.get("Data") or row.get("data"))
            value = as_float(row.get("Total") or row.get("total") or row.get("Valor") or row.get("valor"))
            location = (row.get("Local") or row.get("local") or "").strip() or "Não indicado"
            if not seller or not day:
                continue
            if seller.lower() == "sucata":
                if day == report_date:
                    result["scrap_day"] += value
                result["scrap_month"] += value
            else:
                if day == report_date:
                    result["daily"][seller] += value
                    result["daily_by_location"][location] += value
                result["monthly"][seller] += value
                result["monthly_by_location"][location] += value
            result["daily_series"][day] += value
        return result

    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Use o Excel dos vendedores em formato .xlsx ou .csv.")

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    candidate = None
    for ws in wb.worksheets:
        upper = ws.title.upper()
        if "RESUMO" in upper:
            candidate = ws
            break
    if candidate is None:
        candidate = wb.worksheets[0]

    ws = candidate
    section_headers = []
    for row in range(1, min(ws.max_row, 80) + 1):
        marker = str(ws.cell(row, 1).value or "").strip().upper()
        if marker in {"COIMBRA", "PICOTO"}:
            section_headers.append((row, marker.title()))

    for idx, (header_row, location) in enumerate(section_headers):
        end_row = (section_headers[idx + 1][0] - 1) if idx + 1 < len(section_headers) else ws.max_row
        date_cols: dict[str, int] = {}
        total_col = None
        for col in range(2, ws.max_column + 1):
            value = ws.cell(header_row, col).value
            day = normalize_day(value)
            if day:
                date_cols[day] = col
            elif isinstance(value, str) and "TOTAL" in value.upper():
                total_col = col
        if not date_cols:
            continue
        for row in range(header_row + 1, end_row + 1):
            seller = str(ws.cell(row, 1).value or "").strip()
            if not seller:
                continue
            upper = seller.upper()
            if upper in {"COIMBRA", "PICOTO"} or upper.startswith("PREENCHA"):
                continue
            daily_value = as_float(ws.cell(row, date_cols.get(report_date, -1)).value) if report_date in date_cols else 0.0
            monthly_value = as_float(ws.cell(row, total_col).value) if total_col else sum(as_float(ws.cell(row, col).value) for col in date_cols.values())
            if upper == "SUCATA":
                result["scrap_day"] += daily_value
                result["scrap_month"] += monthly_value
            else:
                result["daily"][seller] += daily_value
                result["monthly"][seller] += monthly_value
                result["daily_by_location"][location] += daily_value
                result["monthly_by_location"][location] += monthly_value
            for day, col in date_cols.items():
                result["daily_series"][day] += as_float(ws.cell(row, col).value)

    # If the workbook has a Dashboard, prefer its explicit daily series when present.
    for dashboard in wb.worksheets:
        if "DASHBOARD" not in dashboard.title.upper():
            continue
        found = {}
        for row in range(1, dashboard.max_row + 1):
            for col in range(1, dashboard.max_column):
                day = normalize_day(dashboard.cell(row, col).value)
                if not day:
                    continue
                value = dashboard.cell(row, col + 1).value
                if isinstance(value, (int, float)):
                    found[day] = float(value)
        if found:
            result["daily_series"] = defaultdict(float, found)
            break

    return result


def derive_operational(payload: dict, report_date: str):
    doc_info = payload.get("documentIntelligence") or {}
    documents = doc_info.get("documents") or []
    selected = [d for d in documents if normalize_day(d.get("date")) == report_date]
    if not selected and documents:
        # Do not silently relabel a different day. Keep empty daily map and let
        # Data Quality flag that the selected day lacks a document-level sample.
        selected = []

    groups = defaultdict(lambda: {"FR_net": 0.0, "FR_total": 0.0, "FT_net": 0.0, "FT_total": 0.0, "NC_net": 0.0, "NC_total": 0.0})
    nc_state = defaultdict(lambda: {"net": 0.0, "total": 0.0})
    scrap = {"net": 0.0, "total": 0.0}

    for doc in selected:
        number = str(doc.get("number") or "")
        dtype = str(doc.get("documentType") or "").upper()
        match_type = re.search(r"\b(FR|FT|NC)\b", dtype + " " + number.upper())
        doc_type = match_type.group(1) if match_type else ("NC" if "NOTA" in dtype else "FT")
        match_series = re.search(r"\b(CUSA|CNOV|PUSA|PNOV|POFI)\b", number.upper())
        series = match_series.group(1) if match_series else "OUTRO"
        net = as_float(doc.get("netAmount"))
        total = as_float(doc.get("totalAmount"))
        entity = str(doc.get("entity") or "")
        if doc_type == "NC":
            net = -abs(net)
            total = -abs(total)
        if "SUCATA" in entity.upper() or "SUCATA" in number.upper():
            scrap["net"] += net
            scrap["total"] += total
            continue
        groups[series][f"{doc_type}_net"] += net
        groups[series][f"{doc_type}_total"] += total
        if doc_type == "NC":
            state = str(doc.get("financialState") or "Desconhecido")
            bucket = "Pendente" if state in {"Pendente", "Vencido"} else "Liquidado"
            nc_state[bucket]["net"] += net
            nc_state[bucket]["total"] += total

    total_net = sum(v[f"{t}_net"] for v in groups.values() for t in ("FR", "FT", "NC")) + scrap["net"]
    total_total = sum(v[f"{t}_total"] for v in groups.values() for t in ("FR", "FT", "NC")) + scrap["total"]
    return {
        "selected_docs": selected,
        "groups": groups,
        "nc_state": nc_state,
        "scrap": scrap,
        "net": total_net,
        "total": total_total,
        "vat": total_total - total_net,
    }


def pstyle(size=10, bold=False, color=INK, leading=None, align=TA_LEFT):
    return ParagraphStyle(
        name=f"seo-{size}-{bold}-{align}",
        parent=getSampleStyleSheet()["BodyText"],
        fontName="Helvetica-Bold" if bold else "Helvetica",
        fontSize=size,
        leading=leading or size * 1.28,
        textColor=color,
        alignment=align,
        spaceAfter=0,
        spaceBefore=0,
    )


def header(c: canvas.Canvas, title: str, subtitle: str, page_no: int):
    c.setFillColor(NAVY)
    c.rect(12 * mm, PAGE_H - 52 * mm, PAGE_W - 24 * mm, 42 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 8)
    c.drawString(25 * mm, PAGE_H - 21 * mm, "SEO CORE")
    c.setFont("Helvetica-Bold", 27)
    c.drawString(25 * mm, PAGE_H - 34 * mm, title)
    c.setFont("Helvetica", 11)
    c.setFillColor(HexColor("#d0d7e0"))
    c.drawString(25 * mm, PAGE_H - 43 * mm, subtitle)
    c.setFillColor(MUTED)
    c.setFont("Helvetica", 7.5)
    c.drawString(14 * mm, 9 * mm, "SEO - Sistema de Eficiência Operacional | Relatório Executivo")
    c.drawRightString(PAGE_W - 14 * mm, 9 * mm, f"Página {page_no}")


def draw_metrics(c, y, metrics):
    x0 = 12 * mm
    width = (PAGE_W - 24 * mm) / len(metrics)
    h = 25 * mm
    for i, (label, value, note) in enumerate(metrics):
        x = x0 + i * width
        c.setStrokeColor(GRID)
        c.setFillColor(colors.white)
        c.rect(x, y - h, width, h, fill=1, stroke=1)
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 7.5)
        c.drawString(x + 5 * mm, y - 6 * mm, label)
        c.setFillColor(INK)
        c.setFont("Helvetica-Bold", 14)
        c.drawString(x + 5 * mm, y - 13 * mm, value)
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 7)
        c.drawString(x + 5 * mm, y - 20 * mm, note[:34])
    return y - h


def draw_paragraph(c, text, x, y, width, style=None):
    style = style or pstyle(9)
    p = Paragraph(text, style)
    w, h = p.wrap(width, PAGE_H)
    p.drawOn(c, x, y - h)
    return y - h


def draw_table(c, data, x, y, widths, font=7.2, header_rows=1, row_heights=None):
    table = Table(data, colWidths=widths, rowHeights=row_heights, repeatRows=header_rows)
    style = [
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), LIGHT),
        ("TEXTCOLOR", (0, 0), (-1, -1), INK),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("FONTNAME", (0, header_rows), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), font),
        ("LEADING", (0, 0), (-1, -1), font + 1.3),
        ("GRID", (0, 0), (-1, -1), 0.35, GRID),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    table.setStyle(TableStyle(style))
    w, h = table.wrap(sum(widths), PAGE_H)
    table.drawOn(c, x, y - h)
    return y - h


def draw_bar_chart(c, items, x, y, w, h, title):
    c.setFillColor(INK)
    c.setFont("Helvetica", 10)
    c.drawCentredString(x + w / 2, y, title)
    y -= 6 * mm
    items = [(name, float(value)) for name, value in items if float(value) > 0][:10]
    if not items:
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 9)
        c.drawString(x, y - 8 * mm, "Sem valores disponíveis para este gráfico.")
        return y - 15 * mm
    max_v = max(v for _, v in items) or 1
    row_h = h / max(1, len(items))
    label_w = 38 * mm
    for i, (name, value) in enumerate(items):
        yy = y - (i + 1) * row_h + row_h * 0.28
        c.setFillColor(INK)
        c.setFont("Helvetica", 7.5)
        c.drawRightString(x + label_w - 2 * mm, yy + 1.5 * mm, name[:25])
        bar_w = (w - label_w - 8 * mm) * value / max_v
        c.setFillColor(BLUE)
        c.rect(x + label_w, yy, bar_w, max(2.8 * mm, row_h * 0.5), fill=1, stroke=0)
    return y - h


def draw_line_chart(c, series, x, y, w, h, title):
    c.setFillColor(INK)
    c.setFont("Helvetica", 10)
    c.drawCentredString(x + w / 2, y, title)
    y -= 7 * mm
    pts = sorted((d, float(v)) for d, v in series.items() if d)
    if len(pts) < 1:
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 9)
        c.drawString(x, y - 8 * mm, "Sem série diária disponível.")
        return y - 15 * mm
    vals = [v for _, v in pts]
    min_v, max_v = min(vals), max(vals)
    span = max(max_v - min_v, max_v * 0.15, 1)
    chart_y = y - h
    c.setStrokeColor(GRID)
    c.rect(x + 13 * mm, chart_y, w - 18 * mm, h, fill=0, stroke=1)
    coords = []
    for i, (day, value) in enumerate(pts):
        px = x + 13 * mm + (w - 18 * mm) * (i / max(1, len(pts) - 1))
        py = chart_y + 6 * mm + (h - 12 * mm) * ((value - min_v) / span)
        coords.append((px, py))
        c.setFillColor(MUTED)
        c.setFont("Helvetica", 6.5)
        c.drawCentredString(px, chart_y - 4 * mm, day[8:10] + "/" + day[5:7])
    c.setStrokeColor(BLUE)
    c.setLineWidth(1.2)
    for a, b in zip(coords, coords[1:]):
        c.line(a[0], a[1], b[0], b[1])
    for px, py in coords:
        c.setFillColor(BLUE)
        c.circle(px, py, 1.8, fill=1, stroke=0)
    return chart_y - 8 * mm


def build_pdf(report_date: str, operational: dict, sellers: dict):
    output = io.BytesIO()
    c = canvas.Canvas(output, pagesize=A4)
    op = derive_operational(operational, report_date)
    report_dt = datetime.fromisoformat(report_date)
    day_label = report_dt.strftime("%d/%m/%Y")

    seller_daily = sorted(sellers["daily"].items(), key=lambda x: x[1], reverse=True)
    seller_month = sorted(sellers["monthly"].items(), key=lambda x: x[1], reverse=True)
    seller_day_total = sum(v for _, v in seller_daily)
    seller_detail_day = seller_day_total + sellers["scrap_day"]
    seller_month_total = sum(v for _, v in seller_month) + sellers["scrap_month"]
    top_day = seller_daily[0] if seller_daily else ("Sem dados", 0.0)
    top_month = seller_month[0] if seller_month else ("Sem dados", 0.0)

    # PAGE 1
    header(c, "Relatório Executivo", f"{day_label} - Faturação diária e acumulado do mês", 1)
    y = PAGE_H - 61 * mm
    y = draw_metrics(c, y, [
        ("Faturação c/ IVA", money(op["total"]), "mapa diário SEO"),
        ("Faturação s/ IVA", money(op["net"]), "mapa diário SEO"),
        ("IVA do dia", money(op["vat"]), "diferença c/ IVA - s/ IVA"),
        ("Documentos", str(len(op["selected_docs"])), "FR + FT + NC processados"),
    ]) - 11 * mm
    c.setFillColor(INK); c.setFont("Helvetica-Bold", 16); c.drawString(16 * mm, y, "Resumo Executivo da IA")
    y -= 8 * mm
    location_totals = defaultdict(float)
    mix_totals = defaultdict(float)
    for series, vals in op["groups"].items():
        total = sum(vals[f"{t}_total"] for t in ("FR", "FT", "NC"))
        loc, group = SERIES_META.get(series, ("Outros", "Outros"))
        location_totals[loc] += total
        mix_totals[group] += total
    op_total_non_scrap = sum(location_totals.values()) or 0.0
    leader_loc = max(location_totals.items(), key=lambda x: x[1])[0] if location_totals else "Sem dados"
    leader_share = (location_totals.get(leader_loc, 0) / op_total_non_scrap * 100) if op_total_non_scrap else 0
    fr_total = sum(v["FR_total"] for v in op["groups"].values())
    nc_total = sum(v["NC_total"] for v in op["groups"].values())
    nc_share = abs(nc_total) / abs(fr_total) * 100 if fr_total else 0
    summary = (
        f"A faturação de {day_label} totaliza <b>{money(op['total'])}</b> com IVA e <b>{money(op['net'])}</b> sem IVA. "
        f"{leader_loc} concentra {pct(leader_share)} do movimento operacional com IVA. "
        f"As notas de crédito somam <b>{money(nc_total)}</b>, equivalentes a {pct(nc_share)} das FR do dia. "
        f"No Excel de vendedores, as vendas atribuídas no dia somam <b>{money(seller_day_total)}</b> e o acumulado registado até ao momento soma <b>{money(seller_month_total)}</b>."
    )
    y = draw_paragraph(c, summary, 16 * mm, y, PAGE_W - 32 * mm, pstyle(10, leading=14)) - 10 * mm
    c.setFont("Helvetica-Bold", 13); c.setFillColor(INK); c.drawString(16 * mm, y, "Destaques do dia")
    y -= 7 * mm
    highlights = [
        ("1", f"{leader_loc} lidera o dia", f"{money(location_totals.get(leader_loc, 0))} c/ IVA ({pct(leader_share)} do movimento operacional)."),
        ("2", "NC com impacto", f"{money(nc_total)} c/ IVA no dia."),
        ("3", "Desempenho comercial", f"{top_day[0]} lidera as vendas atribuídas a vendedores com {money(top_day[1])}."),
    ]
    for num, title, detail in highlights:
        c.setStrokeColor(GRID); c.setFillColor(colors.white); c.rect(15 * mm, y - 18 * mm, PAGE_W - 30 * mm, 16 * mm, fill=1, stroke=1)
        c.setFillColor(INK); c.setFont("Helvetica-Bold", 13); c.drawString(20 * mm, y - 11 * mm, num)
        c.setFont("Helvetica-Bold", 10); c.drawString(34 * mm, y - 7 * mm, title)
        c.setFont("Helvetica", 8); c.setFillColor(MUTED); c.drawString(34 * mm, y - 13 * mm, detail[:105])
        y -= 21 * mm
    c.showPage()

    # PAGE 2 - MAPA DIÁRIO
    header(c, "Mapa Diário SEO", f"{day_label} - Coimbra e Picoto | FR + FT + NC | com e sem IVA", 2)
    y = PAGE_H - 60 * mm
    rows = [["Unidade / Grupo", "Série", "FR c/ IVA", "FR s/ IVA", "FT c/ IVA", "FT s/ IVA", "NC c/ IVA", "NC s/ IVA", "TOTAL c/ IVA", "TOTAL s/ IVA"]]
    ordered = ["CUSA", "CNOV", "PUSA", "PNOV", "POFI"]
    for series in ordered:
        vals = op["groups"].get(series, {})
        loc, group = SERIES_META[series]
        total_total = sum(float(vals.get(f"{t}_total", 0)) for t in ("FR", "FT", "NC"))
        total_net = sum(float(vals.get(f"{t}_net", 0)) for t in ("FR", "FT", "NC"))
        rows.append([
            f"{loc} / {group}", series,
            money(vals.get("FR_total", 0)), money(vals.get("FR_net", 0)),
            money(vals.get("FT_total", 0)), money(vals.get("FT_net", 0)),
            money(vals.get("NC_total", 0)), money(vals.get("NC_net", 0)),
            money(total_total), money(total_net),
        ])
    if op["scrap"]["total"] or sellers["scrap_day"]:
        rows.append(["Sucata - separada", "Separado", "-", "-", "-", "-", "-", "-", money(op["scrap"]["total"] or sellers["scrap_day"]), money(op["scrap"]["net"] or sellers["scrap_day"])])
    rows.append(["TOTAL GERAL", "", "", "", "", "", "", "", money(op["total"]), money(op["net"])])
    widths = [29*mm, 13*mm, 17.5*mm, 17.5*mm, 17.5*mm, 17.5*mm, 17.5*mm, 17.5*mm, 18.5*mm, 18.5*mm]
    y = draw_table(c, rows, 7 * mm, y, widths, font=5.7) - 9 * mm
    c.setFillColor(INK); c.setFont("Helvetica-Bold", 13); c.drawString(12 * mm, y, "Notas de crédito por estado")
    y -= 6 * mm
    nc_rows = [["Estado", "NC c/ IVA", "NC s/ IVA"]]
    for state in ("Liquidado", "Pendente"):
        vals = op["nc_state"].get(state, {})
        nc_rows.append([state, money(vals.get("total", 0)), money(vals.get("net", 0))])
    nc_rows.append(["Total", money(sum(v["total"] for v in op["nc_state"].values())), money(sum(v["net"] for v in op["nc_state"].values()))])
    y = draw_table(c, nc_rows, 12 * mm, y, [58*mm, 60*mm, 60*mm], font=7.5) - 9 * mm
    c.setFillColor(INK); c.setFont("Helvetica-Bold", 13); c.drawString(12 * mm, y, "Composição operacional")
    y -= 6 * mm
    comp = [["Bloco", "c/ IVA", "s/ IVA"]]
    coim_net=coim_total=pic_net=pic_total=used_net=used_total=new_net=new_total=office_net=office_total=0.0
    for series, vals in op["groups"].items():
        loc, group = SERIES_META.get(series, ("Outros", "Outros"))
        ttot=sum(vals[f"{t}_total"] for t in ("FR","FT","NC")); tnet=sum(vals[f"{t}_net"] for t in ("FR","FT","NC"))
        if loc=="Coimbra": coim_total+=ttot; coim_net+=tnet
        if loc=="Picoto": pic_total+=ttot; pic_net+=tnet
        if group=="Usado": used_total+=ttot; used_net+=tnet
        if group=="Novo": new_total+=ttot; new_net+=tnet
        if group=="Oficina": office_total+=ttot; office_net+=tnet
    comp += [
        ["Coimbra", money(coim_total), money(coim_net)],
        ["Picoto operacional", money(pic_total), money(pic_net)],
        ["Sucata", money(op["scrap"]["total"] or sellers["scrap_day"]), money(op["scrap"]["net"] or sellers["scrap_day"])],
        ["Usado (CUSA + PUSA)", money(used_total), money(used_net)],
        ["Novo (CNOV + PNOV)", money(new_total), money(new_net)],
        ["POFI / Oficina", money(office_total), money(office_net)],
    ]
    draw_table(c, comp, 12 * mm, y, [78*mm, 50*mm, 50*mm], font=7.5)
    c.showPage()

    # PAGE 3 - VENDEDORES
    header(c, "Desempenho dos Vendedores", f"{day_label} - ranking diário e acumulado no mês", 3)
    y = PAGE_H - 61 * mm
    y = draw_bar_chart(c, seller_daily, 20 * mm, y, PAGE_W - 40 * mm, 62 * mm, f"Ranking diário de vendedores - {day_label}") - 9 * mm
    total_attr = seller_day_total or 1
    seller_rows = [["Vendedor", "Total líquido do dia", "% das vendas atribuídas"]]
    for name, value in seller_daily[:10]:
        seller_rows.append([name, money(value), pct(value / total_attr * 100)])
    if len(seller_rows) == 1:
        seller_rows.append(["Sem dados para a data", money(0), "0,0%"])
    y = draw_table(c, seller_rows, 12 * mm, y, [100*mm, 48*mm, 42*mm], font=7.7) - 7 * mm
    note = f"Vendas atribuídas a vendedores: {money(seller_day_total)}. Sucata: {money(sellers['scrap_day'])}, apresentada separadamente."
    draw_paragraph(c, note, 14*mm, y, PAGE_W-28*mm, pstyle(8, color=MUTED))
    c.showPage()

    # PAGE 4 - CONSOLIDAÇÃO
    header(c, "Consolidação do Mês", f"Acumulado até {day_label} - faturação líquida e ranking comercial", 4)
    y = PAGE_H - 61 * mm
    month_loc = sellers["monthly_by_location"]
    coim = month_loc.get("Coimbra", 0.0); pic = month_loc.get("Picoto", 0.0)
    month_total = seller_month_total
    active_days = [v for d, v in sellers["daily_series"].items() if d <= report_date and abs(v) > 0.0001]
    avg = sum(active_days)/len(active_days) if active_days else 0
    y = draw_metrics(c, y, [
        ("Total acumulado", money(month_total), "Excel dos vendedores"),
        ("Coimbra", money(coim), pct(coim/month_total*100) if month_total else "0,0%"),
        ("Picoto", money(pic), pct(pic/month_total*100) if month_total else "0,0%"),
        ("Média diária", money(avg), f"{len(active_days)} dias com movimento"),
    ]) - 10 * mm
    series_cut = {d:v for d,v in sellers["daily_series"].items() if d <= report_date}
    y = draw_line_chart(c, series_cut, 20*mm, y, PAGE_W-40*mm, 48*mm, "Faturação líquida diária") - 9 * mm
    draw_bar_chart(c, seller_month, 20*mm, y, PAGE_W-40*mm, 54*mm, "Ranking acumulado de vendedores")
    c.showPage()

    # PAGE 5 - INTELIGÊNCIA
    header(c, "Inteligência Executiva SEO", "Explicação, qualidade e pontos de decisão", 5)
    y = PAGE_H - 61 * mm
    c.setFillColor(INK); c.setFont("Helvetica-Bold", 16); c.drawString(16*mm, y, "Leitura automática dos dados")
    y -= 8*mm
    used_share = used_total / op_total_non_scrap * 100 if op_total_non_scrap else 0
    new_share = new_total / op_total_non_scrap * 100 if op_total_non_scrap else 0
    insights = [
        ("Faturação", f"O dia fecha em {money(op['total'])} com IVA e {money(op['net'])} sem IVA."),
        ("Unidade", f"{leader_loc} responde por {pct(leader_share)} do movimento operacional com IVA do dia."),
        ("Mix", f"Usado concentra {pct(used_share)} e Novo {pct(new_share)} do movimento operacional."),
        ("NC", f"Notas de crédito representam {pct(nc_share)} do valor das FR com IVA."),
        ("Comercial", f"{top_day[0]} lidera o dia com {money(top_day[1])}. {top_month[0]} lidera o acumulado com {money(top_month[1])}."),
    ]
    for label, text in insights:
        c.setStrokeColor(GRID); c.setFillColor(colors.white); c.rect(14*mm, y-16*mm, PAGE_W-28*mm, 14*mm, fill=1, stroke=1)
        c.setFillColor(INK); c.setFont("Helvetica-Bold", 10); c.drawString(19*mm, y-9*mm, label)
        draw_paragraph(c, text, 58*mm, y-5*mm, PAGE_W-76*mm, pstyle(8.4, leading=10.5))
        y -= 19*mm
    y -= 2*mm
    c.setFillColor(INK); c.setFont("Helvetica-Bold", 15); c.drawString(16*mm, y, "Data Quality / Reconciliação")
    y -= 8*mm
    difference = seller_detail_day - op["net"]
    c.setStrokeColor(HexColor("#efb7b0")); c.setFillColor(RED_BG); c.rect(14*mm, y-29*mm, PAGE_W-28*mm, 27*mm, fill=1, stroke=1)
    c.setFillColor(RED); c.setFont("Helvetica-Bold", 10); c.drawString(19*mm, y-10*mm, "ATENÇÃO" if abs(difference) > 0.02 else "VALIDADO")
    quality_text = (
        f"Diferença entre a soma diária do Excel de vendedores + sucata ({money(seller_detail_day)}) e o mapa diário sem IVA ({money(op['net'])}): {money(difference)}. "
        "O SEO mantém os valores das fontes sem os corrigir automaticamente."
    )
    draw_paragraph(c, quality_text, 47*mm, y-5*mm, PAGE_W-65*mm, pstyle(8.3, leading=10.6))
    y -= 37*mm
    c.setFillColor(INK); c.setFont("Helvetica-Bold", 15); c.drawString(16*mm, y, "Recomendação SEO")
    y -= 8*mm
    recommendations = [
        "Rever notas de crédito pendentes antes do fecho oficial.",
        "Validar diferenças entre o mapa diário e o Excel comercial quando superiores à tolerância de 0,02 EUR.",
        "Acompanhar o ritmo diário em relação à média do mês.",
        "Manter sucatas e salvados em blocos independentes no relatório e no cálculo de gestão.",
    ]
    for idx, text in enumerate(recommendations, 1):
        y = draw_paragraph(c, f"{idx}. {text}", 19*mm, y, PAGE_W-38*mm, pstyle(8.5, leading=11)) - 2*mm
    c.showPage()
    c.save()
    output.seek(0)
    return output.getvalue()


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 20 * 1024 * 1024:
                raise ValueError("O pedido do relatório é inválido ou demasiado grande.")
            fields, files = parse_multipart(self.headers, self.rfile.read(length))
            report_date = fields.get("report_date", "").strip()
            operational_json = fields.get("operational_json", "{}")
            if not normalize_day(report_date):
                raise ValueError("Escolha uma data válida para o relatório.")
            if "seller_file" not in files:
                raise ValueError("Anexe o Excel de desempenho dos vendedores.")
            seller_name, seller_bytes = files["seller_file"]
            operational = json.loads(operational_json or "{}")
            sellers = parse_seller_workbook(seller_name, seller_bytes, report_date)
            pdf = build_pdf(report_date, operational, sellers)
            filename = f"Relatorio_Executivo_SEO_{datetime.fromisoformat(report_date).strftime('%d-%m-%Y')}.pdf"
            self.send_response(200)
            self.send_header("Content-Type", "application/pdf")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("Cache-Control", "no-store")
            self.send_header("Content-Length", str(len(pdf)))
            self.end_headers()
            self.wfile.write(pdf)
        except Exception as exc:
            body = json.dumps({"detail": str(exc)}, ensure_ascii=False).encode("utf-8")
            self.send_response(422)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
